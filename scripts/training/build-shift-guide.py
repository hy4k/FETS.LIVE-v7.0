from pathlib import Path
import json, textwrap, os
from fontTools.ttLib import TTFont as FontToolsFont
from fontTools.varLib.instancer import instantiateVariableFont
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT=Path(os.environ.get('FETS_GUIDE_OUTPUT','/tmp/fets-shift-staff-guide'))
FONT=Path(__file__).resolve().parents[2]/'fets-point/src/assets/fonts/manrope-latin-variable.woff2'
for name,weight in [('Regular',450),('Bold',700)]:
    font=FontToolsFont(FONT);font=instantiateVariableFont(font,{'wght':weight},inplace=True);font.flavor=None
    for platform,encoding,language in [(3,1,0x409),(1,0,0)]:
        font['name'].setName(f'ManropeGuide-{name}',6,platform,encoding,language)
    font.save(OUT/f'Manrope-{name}.ttf')
META=json.loads((OUT/'screenshots/metadata.json').read_text())['records']
W,H=1600,900
C={'ink':'2D4538','muted':'62776D','paper':'F8F8F3','white':'FFFFFF','line':'DDE5DC','gold':'BC8B39','staff':'344E40','lead':'94651B','management':'376B88','soft':'E8EEE6','amber':'B98744'}
pdfmetrics.registerFont(TTFont('Guide',str(OUT/'Manrope-Regular.ttf')))
pdfmetrics.registerFont(TTFont('GuideBold',str(OUT/'Manrope-Bold.ttf')))
pdf=canvas.Canvas(str(OUT/'The-Shift-Staff-Guide.pdf'),pagesize=(W,H))
pdf.setTitle('The Shift | Visual staff guide')
pdf.setAuthor('FETS Live')
pdf.setSubject('A pictorial guide to the approved timetable, daily work, checks and handovers')
deck=Presentation();deck.slide_width=Inches(13.333333);deck.slide_height=Inches(7.5)
deck.core_properties.title='The Shift | Visual staff guide'
deck.core_properties.subject='FETS Live staff, lead and management training'
deck.core_properties.author='FETS Live'
slide=None
def color(h): return RGBColor.from_string(C.get(h,h))
def unit(x): return Inches(x/120)
def rect(x,y,w,h,fill=None,stroke=None,r=0,lw=1):
    if fill: pdf.setFillColor(HexColor('#'+C.get(fill,fill)))
    if stroke: pdf.setStrokeColor(HexColor('#'+C.get(stroke,stroke)))
    pdf.setLineWidth(lw)
    if r: pdf.roundRect(x,H-y-h,w,h,r,fill=bool(fill),stroke=bool(stroke))
    else: pdf.rect(x,H-y-h,w,h,fill=bool(fill),stroke=bool(stroke))
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,unit(x),unit(y),unit(w),unit(h))
    if r: shp.adjustments[0]=min(.15,r/min(w,h))
    if fill: shp.fill.solid();shp.fill.fore_color.rgb=color(fill)
    else: shp.fill.background()
    if stroke: shp.line.color.rgb=color(stroke);shp.line.width=Pt(lw*.6)
    else: shp.line.fill.background()
    return shp
def circle(x,y,r,fill='ink'):
    pdf.setFillColor(HexColor('#'+C.get(fill,fill)));pdf.circle(x,H-y,r,stroke=0,fill=1)
    shp=slide.shapes.add_shape(MSO_SHAPE.OVAL,unit(x-r),unit(y-r),unit(2*r),unit(2*r));shp.fill.solid();shp.fill.fore_color.rgb=color(fill);shp.line.fill.background()
def wrapped(s,size,width,bold=False):
    lines=[];font='GuideBold' if bold else 'Guide'
    for para in s.split('\n'):
        line=''
        for word in para.split():
            trial=(line+' '+word).strip()
            if pdfmetrics.stringWidth(trial,font,size)>width and line:lines.append(line);line=word
            else:line=trial
        lines.append(line)
    return lines
def text(s,x,y,size=22,fill='ink',width=1450,bold=False,leading=None):
    lines=wrapped(s,size,width,bold);leading=leading or size*1.32
    pdf.setFillColor(HexColor('#'+C.get(fill,fill)));pdf.setFont('GuideBold' if bold else 'Guide',size)
    for i,l in enumerate(lines):pdf.drawString(x,H-y-size-i*leading,l)
    shp=slide.shapes.add_textbox(unit(x),unit(y),unit(width),unit(len(lines)*leading+8));tf=shp.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,l in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=l;p.font.name='Manrope';p.font.size=Pt(size*.6);p.font.bold=bold;p.font.color.rgb=color(fill);p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=1.02
    return len(lines)*leading
def image(p,x,y,w,h):
    pdf.drawImage(str(p),x,H-y-h,width=w,height=h,mask='auto')
    slide.shapes.add_picture(str(p),unit(x),unit(y),width=unit(w),height=unit(h))
def new(role,page,title,subtitle):
    global slide
    slide=deck.slides.add_slide(deck.slide_layouts[6]);rect(0,0,W,H,'paper')
    text('fets.live',56,34,26,bold=True,width=250);text('THE SHIFT / STAFF GUIDE',232,41,13,fill='muted',width=570,bold=True)
    rect(1266,34,278,34,C.get(role,'ink'),r=16);text('LEADS ONLY' if role == 'lead' else 'STAFF',1284,41,14,'white',width=240,bold=True)
    text(title,56,104,45,bold=True,width=1488)
    text(subtitle,56,174,23,'muted',width=1488)
    rect(56,866,1488,1,'line');text('Timetable edition • fictional training records • 19 October 2026',56,876,12,'muted',width=1300)
    text(f'{page:02}',1496,874,15,'muted',width=48,bold=True)
def screenshot(name,x=480,y=245,w=1064,h=515,marks=True):
    m=META[name];scale=min((w-38)/m['width'],(h-38)/m['height']);iw=m['width']*scale;ih=m['height']*scale
    ix=x+(w-iw)/2;iy=y+(h-ih)/2
    rect(x,y,w,h,'white','line',r=18)
    image(OUT/'screenshots'/f'{name}.png',ix,iy,iw,ih)
    if marks:
        for p in m['points']:
            px=ix+p['x']*scale;py=iy+p['y']*scale;pw=p['w']*scale;ph=p['h']*scale
            rect(px-3,py-3,pw+6,ph+6,stroke='amber',r=4,lw=2)
            bx=max(x+18,px-6);by=max(y+18,py-6)
            circle(bx,by,16,'ink');text(str(p['n']),bx-5,by-12,16,'white',width=22,bold=True)
    return ix,iy,iw,ih
def callouts(items,role):
    y=260;fill=role
    title_size=24 if len(items)<4 else 22
    detail_size=21 if len(items)<4 else 19
    gap=37 if len(items)<4 else 24
    for n,(title,detail) in enumerate(items,1):
        circle(77,y+17,18,fill);text(str(n),71,y+3,19,'white',width=30,bold=True)
        ht=text(title,110,y,title_size,bold=True,width=330)
        hd=text(detail,110,y+ht+7,detail_size,'muted',width=320)
        y+=ht+hd+gap
    assert y<779,('Callouts overflow',items,y)
def takeaway(s):
    rect(56,790,1488,60,'soft',r=13);text('KEY POINT',77,813,12,'staff',width=134,bold=True);text(s,225,805,22,width=1290)

SLIDES=[
 ('staff','Open your timetable','Use The Shift from the main menu.','open',[
  ('Open The Shift','Select The Shift, beside Actionables.'),
  ('Check centre and date','Make sure you are looking at your centre and the correct working day.'),
  ('Choose My shift','This is your personal timetable. Team shows the centre’s responsibilities.')],
  'Use this guide with the timetable edition shown here. The centre selector is in the main application.'),
 ('staff','Read the whole day in time order','Posts, breaks and work deadlines share one timetable.','timetable',[
  ('Jump to now','Press the clock to move to your current post.'),
  ('Follow the green row','This is your current responsibility. Read its instructions and end time.'),
  ('Notice BY times','These are work deadlines. Open the duty to see the expected result.')],
  'A post time tells you where to work. A BY time tells you when an assigned result is due.'),
 ('staff','Read the result, then start','Open a duty row to see its instructions and controls.','start',[
  ('Read Done means','Understand the result required. Open Instructions for the method or checklist.'),
  ('Select Start duty','Do this when you actually begin the work.'),
  ('Use Need help','Ask the lead if the result, equipment or next step is unclear.')],
  'Start duty records progress. It does not mark the work complete.'),
 ('staff','Ask for help clearly','Open the duty, then select Need help.','help',[
  ('Explain the problem','Say what is stopping you and what would help.'),
  ('Choose Help needed','Select the most useful reason, such as equipment or clearer instructions.'),
  ('Send the request','Select Ask the lead for help. Use Resume work after the blocker is resolved.')],
  'The duty remains assigned to you unless the lead reassigns it.'),
 ('staff','Finish with a real result','When the work is complete, select Finish duty.','finish',[
  ('Describe what you completed','State what you checked or delivered. Include a useful reference and any unresolved issue.'),
  ('Send for review','An independent reviewer will check the result against the agreed standard.')],
  'Write the actual outcome. Do not copy the example wording unless it describes your work.'),
 ('staff','Know what Ready for review means','Submitting a result is followed by independent verification.','submitted',[
  ('Read the status','Ready for review means the result has been submitted, not verified.'),
  ('Check your recorded result','Make sure the description is accurate and complete.'),
  ('Wait for the reviewer','The status becomes Verified after an independent check. Open the duty for feedback.')],
  'Reopen my work is for a correction or further work, not a way to verify your own result.'),
 ('staff','Respond to returned work','A reviewer may ask for a specific correction.','returned',[
  ('Read Reviewer feedback','Check exactly what needs to change. Ask for help if it is unclear.'),
  ('Correct, then resubmit','Complete the correction, select Finish duty and record the updated result.')],
  'A returned result stays open until the required correction and independent review are complete.'),
 ('staff','Record the lab walk after doing it','Lab walks follow the 10-minute interval.','walk',[
  ('Follow your assigned post','Use the current responsibility and published cover arrangement.'),
  ('Select Record lab walk','Complete the walk first, then record it when the button is available.')],
  'Current process: record in The Shift. FETS Online integration will be handled in a separate stage.'),
 ('staff','Record the DVR check after doing it','DVR checks follow the 6-minute interval.','dvr',[
  ('Check who owns the post','Read your current responsibility, including any named break cover.'),
  ('Select Record DVR check','Inspect the DVR first, then record the check using the available button.')],
  'Do not enter checks that did not happen. Earlier missing checks cannot be backfilled here.'),
 ('staff','Take the break with clear cover','The timetable shows the agreed break and responsible cover.','break',[
  ('Check the start and return time','Follow the published interval and tell the lead if the plan cannot be followed.'),
  ('Confirm the named cover','Make sure the named colleague understands the post and its checks before you leave.')],
  'During break cover, the named cover owns the post and its checks. Arrange gaps with the lead.'),
 ('staff','See who is responsible across the centre','Open Team to see the same day for everyone.','team',[
  ('Read the named owners','Each post shows who is responsible for that exact interval.'),
  ('Check the break note','The break and effective cover are shown together. Contact the lead about any mismatch.')],
  'Use the published assignment. Do not assume the same person keeps a post for the entire day.'),
 ('staff','Leave a useful post handover','Open Handover records below the timetable, then select the post.','handover',[
  ('Choose the correct interval','Check the post and its times.'),
  ('Describe the outcome','State what is complete and what is still open, including the next owner or follow-up.'),
  ('Submit handover','Available for a post you covered after its rotation ends. The lead can review the result.')],
  'Report missing checks honestly. A handover does not create or replace the individual check records.'),
 ('staff','Read feedback and add your perspective','Open My feedback for your private feedback and follow-up.','feedback',[
  ('Add relevant context','Explain circumstances, workload or support that affected the work.'),
  ('Save my response','Your response stays with the feedback so management can follow up.')],
  'Recognition, support and training should be based on actual work and context, not counts alone.'),
 ('lead','Review the result independently','Open Team, find the submitted duty and open it.','review',[
  ('Read the recorded result','Compare it with Done means and check the actual work.'),
  ('Verify result','Use this when the required standard has been met.'),
  ('Return with feedback','Describe the correction needed, then Save decision. The owner can resubmit.')],
  'You cannot verify your own completed duty. The lead’s regular duties still need an independent review.'),
 ('lead','Check the day report before sending','Open Report after the day’s rota is published.','report',[
  ('Use today’s records','Use the records as a starting point, then check the summary.'),
  ('Carry into tomorrow','List unresolved work with a clear next owner and target. Write None if nothing remains.'),
  ('Send to super admin','Available from 17:00 India time. Complete reviews and corrections first.')],
  'Sending the report closes that day’s shift records. Check the report carefully before sending.'),
 ('staff','Use the same timetable on your phone','The same workflow is available on the smaller screen.','phone',[
  ('Tap the clock','Move straight to your current responsibility.'),
  ('Tap a duty row','Read the expected result, then start, ask for help or finish the work.')],
  'Check your day → do the work → record the real result → respond to independent review.'),
]

for i,(role,title,subtitle,name,items,remember) in enumerate(SLIDES,1):
    assert len(items)==len(META[name]['points']),('Numbered instructions must match screenshot highlights',name)
    new(role,i,title,subtitle);callouts(items,role);screenshot(name);takeaway(remember);pdf.showPage()
pdf.save();deck.save(OUT/'The-Shift-Staff-Guide.pptx')
manifest={'pages':len(SLIDES),'files':['The-Shift-Staff-Guide.pdf','The-Shift-Staff-Guide.pptx'],'screenshots':list(META),'training_data':True,'edition':'Approved timetable preview, 19 October 2026','online_integration':'Future stage; current screenshots show recording in The Shift','steps':[{'page':i,'title':row[1],'role':row[0],'screenshot':row[3]} for i,row in enumerate(SLIDES,1)]}
(OUT/'guide-manifest.json').write_text(json.dumps(manifest,indent=2))
print(json.dumps(manifest))
